from docx import Document
from pathlib import Path
from docx.oxml.ns import qn

src = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Informacion_Preliminar_Establecimiento_Requerimientos_Procesos_Aplicativo_Inventarios.docx')
dst = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Informacion_Preliminar_Establecimiento_Requerimientos_Procesos_Aplicativo_Inventarios.md')
doc = Document(src)

def esc(s):
    return s.replace('|', '\\|').replace('\n', ' ')

def table_md(table):
    rows = [[esc(c.text.strip()) for c in row.cells] for row in table.rows]
    if not rows: return ''
    out = ['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * len(rows[0])) + ' |']
    out += ['| ' + ' | '.join(r) + ' |' for r in rows[1:]]
    return '\n'.join(out)

paras = {p._p: p for p in doc.paragraphs}
tables = {t._tbl: t for t in doc.tables}
lines = []
for el in doc._body._element.iterchildren():
    if el.tag == qn('w:p'):
        p = paras.get(el)
        if not p or not p.text.strip():
            continue
        text = p.text.strip()
        style = p.style.name if p.style else ''
        if style == 'Title':
            lines += ['# ' + text, '']
        elif style.startswith('Heading 1'):
            lines += ['## ' + text, '']
        elif style.startswith('Heading 2'):
            lines += ['### ' + text, '']
        elif style.startswith('List Bullet'):
            lines += ['- ' + text, '']
        else:
            lines += [text, '']
    elif el.tag == qn('w:tbl'):
        t = tables.get(el)
        if t:
            lines += [table_md(t), '']

dst.write_text('\n'.join(lines).rstrip() + '\n', encoding='utf-8')
print(dst)
