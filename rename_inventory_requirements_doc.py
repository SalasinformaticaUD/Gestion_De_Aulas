from docx import Document
from pathlib import Path

old = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Levantamiento_Preliminar_Inventarios_Salas_Informatica.docx')
new = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Informacion_Preliminar_Establecimiento_Requerimientos_Procesos_Aplicativo_Inventarios.docx')
doc = Document(old)
title = 'Información preliminar para el establecimiento de requerimientos y procesos del aplicativo de inventarios'
if doc.paragraphs:
    p = doc.paragraphs[0]
    for r in p.runs: r.text = ''
    if p.runs: p.runs[0].text = title
    else: p.add_run(title)
tmp = new.with_name(new.stem + '_tmp.docx')
doc.save(tmp)
tmp.replace(new)
if old.exists(): old.unlink()
print(new)
