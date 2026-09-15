from docx import Document
from pathlib import Path

path = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Levantamiento_Preliminar_Inventarios_Salas_Informatica.docx')
doc = Document(path)

replacements = {
    'Cada activo tendr� un identificador interno �nico y permanente denominado id_activo. Este identificador ser� la llave t�cnica del aplicativo, aun cuando la placa institucional o el serial cambien, se corrijan o se repitan entre activos independientes.':
    'Los activos se identifican en las fuentes mediante placas institucionales, n�meros de serie, c�digos de inventario y otros campos descriptivos. La relaci�n entre estos identificadores debe conservarse en los formatos institucionales y validarse durante la conciliaci�n.',
    'La ficha central del activo ser� la fuente operativa de los datos compartidos. Desde ella se alimentar�n simult�neamente la hoja de vida y el formato LAI. El id_activo no sustituir� la placa ni el serial en los formatos; estos conservar�n los campos visibles y el dise�o institucional vigente.':
    'La ficha central del activo ser� la fuente operativa de los datos compartidos. Desde ella se alimentar�n simult�neamente la hoja de vida y el formato LAI. La ficha no sustituir� la placa ni el serial en los formatos; estos conservar�n los campos visibles y el dise�o institucional vigente.'
}

def replace_in_paragraph(p):
    if p.text in replacements:
        # preserve paragraph style while replacing all runs
        for r in p.runs:
            r.text = ''
        if p.runs:
            p.runs[0].text = replacements[p.text]
        else:
            p.add_run(replacements[p.text])

for p in doc.paragraphs:
    if 'id_activo' in p.text:
        if 'llave' in p.text or 'identificador interno' in p.text:
            new = 'Los activos se identifican en las fuentes mediante placas institucionales, n�meros de serie, c�digos de inventario y otros campos descriptivos. La relaci�n entre estos identificadores debe conservarse en los formatos institucionales y validarse durante la conciliaci�n.'
        else:
            new = 'La ficha central del activo ser� la fuente operativa de los datos compartidos. Desde ella se alimentar�n simult�neamente la hoja de vida y el formato LAI. La ficha no sustituir� la placa ni el serial en los formatos; estos conservar�n los campos visibles y el dise�o institucional vigente.'
        for r in p.runs: r.text=''
        if p.runs: p.runs[0].text=new
        else: p.add_run(new)

table_repls = {
    'Se relaciona con id_activo y se desambigua por tipo, serial y contexto.':
        'Se desambigua por tipo, serial, placa y contexto.',
    'Se registra el serial y se mantiene el id_activo interno.':
        'Se registra el serial como referencia del componente o proceso.',
    'El proceso queda asociado al c�digo interno y al id_activo.':
        'El proceso queda asociado al c�digo interno cuando este haya sido generado por almac�n.'
}
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            old_text = cell.text
            if old_text in table_repls or 'id_activo' in old_text:
                new_text = table_repls.get(old_text, 'El proceso queda asociado al c�digo interno cuando este haya sido generado por almac�n.')
                for p in cell.paragraphs:
                    if p.text == old_text:
                        for r in p.runs: r.text=''
                        if p.runs: p.runs[0].text = new_text
                        else: p.add_run(new_text)

doc.save(path)
print('updated')
