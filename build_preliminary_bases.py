from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output')
OUT.mkdir(exist_ok=True)

md = '''# Bases documentales para definir los macroprocesos del aplicativo

## Propósito y alcance

Este documento consolida únicamente información observada en los archivos y carpetas suministrados. Su finalidad es servir como base para definir los macroprocesos del aplicativo de inventarios, equipos, salas, mantenimientos, traslados e insumos. No incorpora reglas de diseño, claves técnicas, decisiones de arquitectura ni procedimientos que no estén explícitamente respaldados por las fuentes.

## Fuentes revisadas

| Fuente | Contenido observado |
|---|---|
| `Hojas_Vida_Salas_Informatica.xlsm` | Libro principal con hojas de vida, mantenimientos, base de datos, inventario de profesores, formatos y lista de imágenes. |
| `Hojas_Vida_Salas_Informatica_1.csv` a `_8.csv` | Exportaciones de hojas activas del libro principal. |
| `Hojas_Vida_Salas_Informatica.pdf` | Representación del formato contenido en la hoja 2, correspondiente a la Hoja de Vida. |
| `inventario insumos.xlsx` | Catálogo de insumos, estados, historial, trabajadores, ubicaciones y empaques. |
| Carpeta `C:\\Users\\MONITORES\\Downloads\\INVENTARIOS` | Formatos LAI, recepción técnica, transferencias, servicio técnico, placas, inventarios históricos, contenido de discos y soportes documentales. |

## Hallazgos documentales

### 1. Activos y hojas de vida

- El libro principal contiene una base de datos de activos en la hoja `Base de Datos`, con aproximadamente 1.154 registros y campos de identificación, descripción, ubicación y características.
- La hoja `Hoja de vida equipos` contiene un formato institucional estructurado, con áreas combinadas, imágenes, fórmulas y zona de impresión definida. El PDF suministrado corresponde a este formato.
- Se observan equipos informáticos y periféricos, entre ellos CPU, monitores, teclados, mouse, proyectores o video beams, pantallas interactivas, tablets y portátiles.
- Los inventarios y formatos también documentan elementos de sala y mobiliario, como sillas, mesas, parlantes y otros bienes.
- En las fuentes aparecen placas institucionales y números de serie como datos de identificación. La forma en que estos identificadores se relacionan debe conservarse en los formatos institucionales.

### 2. Levantamiento de salas y formato LAI

- Los formatos LAI corresponden al código `GIF-PR-003-FR-004`, versión 1.1.
- El levantamiento se organiza por sala, campaña o periodo, e incluye responsables, revisores y fechas.
- Se registran verificaciones física, ARKA y ERP, además de estados como Activo, Obsoleto, Dañado y Mantenimiento.
- La hoja de mobiliario incluye ítems de sala, identificadores locales y números institucionales.
- El formato LAI es un documento institucional y debe conservar su estructura y contenido oficial.

### 3. Recepción técnica y alta documental

- El formato de recepción `GSIT-PR-002-FR-007`, versión 2, registra orden de compra, proveedor, especificaciones técnicas, seriales de CPU y monitor, lista de comprobación funcional, garantía y firmas.
- En la documentación revisada se encuentra un acta de recepción de 183 computadores correspondiente a 2024.
- La recepción técnica constituye un punto documental previo a la incorporación o actualización de información de equipos.

### 4. Mantenimientos y servicio técnico

- La hoja `Mantenimientos` contiene aproximadamente 3.242 registros y cerca de 900 identificadores únicos.
- Se identificaron 43 registros de mantenimiento sin correspondencia clara en la base de activos revisada.
- El formato de servicio `SG-FR34` incluye número de caso, falla reportada, estado físico, diagnóstico, pruebas, solución, repuestos, observaciones y firmas del cliente y del técnico.
- La documentación registra mantenimientos preventivos y correctivos, así como intervenciones asociadas a componentes.

### 5. Placas, componentes y relación entre equipos

- Los archivos `SN_SALAS_PLAQUETA`, `PCS_NUEVOS`, `PCS_ANTIGUOS` e inventarios históricos contienen relaciones entre torres, monitores y otros componentes.
- CPU, monitor, mouse, teclado y periféricos aparecen como elementos identificables en los inventarios y formatos.
- Algunas fuentes contienen placas repetidas en contextos distintos; el significado operativo de estas repeticiones debe validarse antes de convertirlas en restricciones del sistema.

### 6. Traslados, entregas y asignaciones

- El formato de transferencia `GIF-PR-005-FR-008` registra quien entrega, quien recibe, fechas, elementos transferidos, origen, destino y observaciones.
- Los documentos de entrega y recepción contienen información de responsables y soportes de la operación.
- La documentación revisada contempla trabajadores, responsables de sala y personal técnico o administrativo.

### 7. Insumos

- `inventario insumos.xlsx` contiene las hojas `empaques`, `estados`, `historial`, `insumos`, `trabajadores` y `ubicaciones`.
- El catálogo contiene aproximadamente 906 registros de insumos; se observan existencias actuales, cantidades iniciales, estado, ubicación y datos de empaque.
- El historial contiene aproximadamente 1.529 eventos con tipos como Nuevo ingreso, Eliminar, Editar, Salida, Entrada, Entradas/Salidas y Reactivar.
- Gran parte del detalle de los movimientos está vacío o expresado en texto libre; por ello, el historial no permite reconstruir con certeza todos los saldos únicamente a partir de sus textos.
- Se observan trabajadores y ubicaciones como catálogos relacionados con la gestión de insumos.
- En otra versión revisada de inventario se identifican movimientos con columnas de cantidad más estructuradas.

### 8. Imágenes, discos, copias y soportes

- `Lista Imagenes` y `Contenido Discos.xlsx` documentan imágenes, respaldos, plantillas, software y contenido asociado a equipos o salas.
- La documentación permite identificar la existencia de estos soportes, pero no define por sí sola si todos deben formar parte del alcance funcional del aplicativo.

## Macroprocesos que pueden delimitarse con la evidencia disponible

1. **Inventario y caracterización de activos y salas.** Consolidación de equipos, periféricos, mobiliario, placas, seriales, ubicaciones, estados y características.
2. **Gestión de hojas de vida institucionales.** Consulta y diligenciamiento del formato oficial de Hoja de Vida, con sus imágenes, campos y antecedentes.
3. **Levantamiento y conciliación LAI.** Registro de campañas por sala y comparación de inventario físico, ARKA y ERP.
4. **Recepción técnica de equipos.** Registro de adquisiciones, proveedores, órdenes, especificaciones, pruebas, garantías y firmas.
5. **Traslados, entregas y asignaciones.** Documentación de movimientos de bienes entre responsables, salas, sedes, almacén y centro de acopio.
6. **Mantenimiento y soporte técnico.** Gestión de casos, diagnósticos, pruebas, soluciones, repuestos, garantías y cierre.
7. **Gestión de insumos y movimientos.** Catálogo, existencias, entradas, salidas, ajustes, estados, ubicaciones, trabajadores e historial.
8. **Soportes documentales y trazabilidad.** Gestión de formatos, firmas, imágenes, respaldos, observaciones y consultas históricas.

## Límites de lo que puede afirmarse actualmente

Los archivos permiten delimitar dominios, documentos, datos y relaciones observadas. No permiten establecer todavía, de manera concluyente, los permisos por rol, aprobaciones, estados de transición, reglas de unicidad, política de integración futura con ERP, procedimiento de migración, retención documental ni alcance definitivo de imágenes, copias y software.

## Base para la siguiente etapa

La siguiente etapa debe validar cada macroproceso con sus responsables y convertir los hallazgos en: entradas y salidas, actividades, documentos utilizados, datos obligatorios, excepciones, responsables y criterios de cierre. Las definiciones técnicas o reglas de negocio deberán registrarse separadamente como decisiones aprobadas, no como hallazgos.
'''

(OUT / 'Bases_Documentales_Macroprocesos_Inventarios.md').write_text(md, encoding='utf-8')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def mark_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7); sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(22); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
for s in ('Heading 1','Heading 2'):
    styles[s].font.name='Aptos'; styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(14); styles['Heading 1'].font.bold=True
styles['Heading 2'].font.size=Pt(11.5); styles['Heading 2'].font.bold=True

title = doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.LEFT; title.add_run('Bases documentales para definir los macroprocesos del aplicativo')
p=doc.add_paragraph(); p.add_run('Inventarios de salas de informática, equipos, mantenimientos e insumos').italic=True
p=doc.add_paragraph(); p.add_run('Propósito. ').bold=True; p.add_run('Consolidar los hallazgos verificables en los archivos revisados para delimitar los macroprocesos del aplicativo, sin convertir decisiones de diseño pendientes en reglas de negocio.')

def h(text, level=1): doc.add_heading(text, level=level)
def para(text): doc.add_paragraph(text)
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')

h('Fuentes revisadas')
table=doc.add_table(rows=1, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for i,t in enumerate(['Fuente','Contenido observado']):
    c=table.rows[0].cells[i]; c.text=t; set_cell_shading(c,'1F4E78'); set_cell_border(c)
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
mark_header(table.rows[0])
sources=[('Hojas_Vida_Salas_Informatica.xlsm y CSV 1 a 8','Hojas de vida, mantenimientos, base de datos, inventario de profesores, formatos y lista de imágenes.'),('Hojas_Vida_Salas_Informatica.pdf','Representación del formato de Hoja de Vida contenido en la hoja 2.'),('inventario insumos.xlsx','Catálogo, estados, historial, trabajadores, ubicaciones y empaques.'),('Carpeta INVENTARIOS','Formatos LAI, recepción técnica, transferencias, servicio, placas, inventarios históricos y soportes.')]
for a,b in sources:
    cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
    for c in cells: set_cell_border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

sections=[
('Activos y hojas de vida',['La hoja Base de Datos contiene aproximadamente 1.154 registros con identificación, descripción, ubicación y características.','La hoja Hoja de vida equipos es un formato institucional estructurado, con áreas combinadas, imágenes, fórmulas y zona de impresión.','Se observan CPU, monitores, teclados, mouse, proyectores o video beams, pantallas interactivas, tablets, portátiles y mobiliario.','Placas institucionales y números de serie aparecen como datos de identificación.']),
('Levantamiento de salas y LAI',['Los formatos LAI corresponden a GIF-PR-003-FR-004, versión 1.1.','Se organizan por sala, campaña o periodo, con responsables, revisores y fechas.','Incluyen verificaciones física, ARKA y ERP; se observan estados Activo, Obsoleto, Dañado y Mantenimiento.','El formato incluye mobiliario, identificadores locales y números institucionales; su estructura oficial debe conservarse.']),
('Recepción técnica',['GSIT-PR-002-FR-007 v2 registra orden de compra, proveedor, especificaciones, seriales, pruebas funcionales, garantía y firmas.','Se encontró un acta de recepción de 183 computadores de 2024.']),
('Mantenimientos y servicio',['Mantenimientos contiene aproximadamente 3.242 registros y cerca de 900 identificadores únicos.','Se identificaron 43 registros sin correspondencia clara en la base de activos revisada.','SG-FR34 registra caso, falla, estado físico, diagnóstico, pruebas, solución, repuestos, observaciones y firmas.']),
('Placas y componentes',['SN_SALAS_PLAQUETA, PCS_NUEVOS, PCS_ANTIGUOS e inventarios históricos relacionan torres, monitores y otros componentes.','CPU, monitor, mouse, teclado y periféricos aparecen como elementos identificables.','Se observan placas repetidas en contextos distintos; los archivos no definen por sí solos una restricción de unicidad.']),
('Traslados y asignaciones',['GIF-PR-005-FR-008 registra quien entrega, quien recibe, fechas, elementos, origen, destino y observaciones.','Los documentos contemplan trabajadores, responsables de sala y personal técnico o administrativo.']),
('Insumos',['inventario insumos.xlsx contiene empaques, estados, historial, insumos, trabajadores y ubicaciones.','El catálogo contiene aproximadamente 906 registros; se observan existencias, cantidades iniciales, estados, ubicaciones y empaques.','El historial contiene aproximadamente 1.529 eventos. Gran parte del detalle está vacío o en texto libre, por lo que no permite reconstruir con certeza todos los saldos.','Otra versión revisada contiene movimientos con cantidades más estructuradas.']),
('Imágenes y soportes',['Lista Imagenes y Contenido Discos.xlsx documentan imágenes, respaldos, plantillas, software y contenido asociado a equipos o salas.','Los archivos prueban la existencia de estos soportes, pero no determinan por sí solos que todos formen parte del alcance funcional.'])]
for name, items in sections:
    h(name); bullets(items)

h('Macroprocesos delimitables con la evidencia')
macro=[('Inventario y caracterización','Equipos, periféricos, mobiliario, placas, seriales, ubicaciones, estados y características.'),('Hojas de vida institucionales','Consulta y diligenciamiento del formato oficial y sus antecedentes.'),('Levantamiento y conciliación LAI','Campañas por sala y comparación física, ARKA y ERP.'),('Recepción técnica','Adquisiciones, proveedores, órdenes, especificaciones, pruebas, garantías y firmas.'),('Traslados, entregas y asignaciones','Movimientos entre responsables, salas, sedes, almacén y centro de acopio.'),('Mantenimiento y soporte técnico','Casos, diagnósticos, pruebas, soluciones, repuestos, garantías y cierres.'),('Gestión de insumos','Catálogo, existencias, entradas, salidas, ajustes, estados, ubicaciones e historial.'),('Soportes y trazabilidad','Formatos, firmas, imágenes, respaldos, observaciones y consultas históricas.')]
t=doc.add_table(rows=1, cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,x in enumerate(['Macroproceso','Base documental']):
    c=t.rows[0].cells[i]; c.text=x; set_cell_shading(c,'1F4E78'); set_cell_border(c)
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
mark_header(t.rows[0])
for a,b in macro:
    cs=t.add_row().cells; cs[0].text=a; cs[1].text=b
    for c in cs: set_cell_border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

h('Límites de la evidencia')
para('Los archivos permiten delimitar dominios, documentos, datos y relaciones observadas. No permiten establecer de forma concluyente permisos por rol, aprobaciones, estados de transición, reglas de unicidad, política de integración futura con ERP, procedimiento de migración, retención documental ni alcance definitivo de imágenes, copias y software.')
h('Siguiente etapa')
para('Validar cada macroproceso con sus responsables y convertir los hallazgos en entradas, salidas, actividades, documentos, datos obligatorios, excepciones, responsables y criterios de cierre. Las decisiones técnicas o reglas de negocio deberán registrarse por separado como decisiones aprobadas.')

doc.save(OUT / 'Bases_Documentales_Macroprocesos_Inventarios.docx')
print('created')
