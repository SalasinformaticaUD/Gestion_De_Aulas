from docx import Document
from pathlib import Path

path = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Levantamiento_Preliminar_Inventarios_Salas_Informatica.docx')
doc = Document(path)
if doc.paragraphs:
    p = doc.paragraphs[0]
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = 'Informaci\u00f3n preliminar para los macroprocesos del aplicativo'
    else:
        p.add_run('Informaci\u00f3n preliminar para los macroprocesos del aplicativo')
tmp = path.with_name(path.stem + '_tmp.docx')
doc.save(tmp)
tmp.replace(path)
print('updated')
