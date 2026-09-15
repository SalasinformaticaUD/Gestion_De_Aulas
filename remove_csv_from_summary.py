from docx import Document
from pathlib import Path

path = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Levantamiento_Preliminar_Inventarios_Salas_Informatica.docx')
doc = Document(path)
old_start = 'La evidencia demuestra que el archivo XLSM'
for p in doc.paragraphs:
    if p.text.startswith(old_start):
        new = 'La evidencia demuestra que el archivo XLSM es la fuente de an�lisis adecuada, porque conserva la estructura del libro, las im�genes, las f�rmulas, las celdas combinadas, la configuraci�n de impresi�n y los componentes necesarios para comprender el proceso actual.'
        for r in p.runs: r.text = ''
        if p.runs: p.runs[0].text = new
        else: p.add_run(new)
        break
tmp = path.with_name(path.stem + '_tmp.docx')
doc.save(tmp)
tmp.replace(path)
print('updated')
