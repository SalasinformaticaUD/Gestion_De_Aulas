from docx import Document
from pathlib import Path

path = Path(r'C:\Users\MONITORES\Documents\Software Monitorias\output\Informacion_Preliminar_Establecimiento_Requerimientos_Procesos_Aplicativo_Inventarios.docx')
doc = Document(path)
repls = {'an�lisis':'análisis','im�genes':'imágenes','f�rmulas':'fórmulas','configuraci�n':'configuración','impresi�n':'impresión','conciliaci�n':'conciliación','n�meros':'números','c�digos':'códigos','relaci�n':'relación','c�digo':'código','almac�n':'almacén'}
def fix(text):
    for a,b in repls.items(): text=text.replace(a,b)
    return text
for p in doc.paragraphs:
    if '�' in p.text:
        new=fix(p.text)
        for r in p.runs: r.text=''
        if p.runs: p.runs[0].text=new
        else: p.add_run(new)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if '�' in cell.text:
                old=cell.text; new=fix(old)
                for p in cell.paragraphs:
                    if p.text==old:
                        for r in p.runs:r.text=''
                        if p.runs:p.runs[0].text=new
                        else:p.add_run(new)
tmp=path.with_name(path.stem+'_tmp.docx'); doc.save(tmp); tmp.replace(path)
print('updated')
